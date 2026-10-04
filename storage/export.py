"""导出 Excel：逐份对照表 + 逐点判定明细。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from storage.db import ROOT

EXPORT_DIR = ROOT / "导出"

HEADER_FILL = PatternFill("solid", fgColor="EFE9DA")
RED = PatternFill("solid", fgColor="FBEFEF")
YELLOW = PatternFill("solid", fgColor="FDF6E3")
GREEN = PatternFill("solid", fgColor="F0F5E8")


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def export_excel(paper: dict) -> Path:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    question = paper["question"]
    students = paper["students"]
    ai = paper["ai"]
    teacher = paper["teacher"]

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    path = EXPORT_DIR / f"批改记录_第{paper['paper_id']}份_{stamp}.xlsx"

    wb = Workbook()

    # ---- 表一：总览 ----
    ws = wb.active
    ws.title = "总览"
    ws.append(["题目", question.stem])
    ws.append(["主题", question.topic])
    ws.append(["满分", question.max_score])
    ws.append(["生成时间", paper.get("created_at", "")])
    ws.append(["AI 判定次数/份", paper.get("repeats", "")])

    # 耗时统计（2026-10-04 新增）。老存档没这些数据，就不写这几行。
    # ⚠ 延迟导入：repository 也依赖 db，模块级导入容易形成循环依赖
    from core.timing import fmt_duration
    from storage.repository import load_timing

    _t = load_timing(paper["paper_id"])
    if _t.get("t_total") is not None:
        ws.append(["出题耗时", fmt_duration(_t.get("t_question"))])
        ws.append([f"生成答卷耗时（{len(students)} 份）",
                   fmt_duration(_t.get("t_answer"))])
        ws.append([f"AI 阅卷耗时（{len(students)} 份）",
                   fmt_duration(_t.get("t_grade"))])
        ws.append(["总耗时", fmt_duration(_t.get("t_total"))])
        if _t.get("api_calls"):
            ws.append(["API 调用次数", _t["api_calls"]])
    ws.append([])
    for p in question.points:
        ws.append([f"采分点{p.seq}", p.text, p.score])
    for i in range(1, 7):
        ws.cell(row=i, column=1).font = Font(bold=True)

    # ---- 表二：逐份对照 ----
    ws2 = wb.create_sheet("逐份对照")
    ws2.append(
        ["序号", "水平档位", "AI分", "AI置信度", "教师分", "差值", "低置信", "学生答案"]
    )
    for s in students:
        a = ai.get(s.seq)
        t = teacher.get(s.seq)
        diff = (t - a.total) if (t is not None and a is not None) else None
        ws2.append(
            [
                s.seq,
                s.ability,
                a.total if a else "",
                a.confidence if a else "",
                t if t is not None else "",
                diff if diff is not None else "",
                "是" if (a and a.low_confidence) else "",
                s.answer,
            ]
        )
        row = ws2.max_row
        if diff is not None:
            if abs(diff) < 1e-6:
                fill = GREEN
            elif abs(diff) <= 1:
                fill = YELLOW
            else:
                fill = RED
            for col in range(1, 9):
                ws2.cell(row=row, column=col).fill = fill
        ws2.cell(row=row, column=8).alignment = Alignment(wrap_text=True, vertical="top")
    ws2.column_dimensions["H"].width = 80
    ws2.column_dimensions["A"].width = 6
    _style_header(ws2)

    # ---- 表三：逐点判定明细 ----
    ws3 = wb.create_sheet("逐点判定")
    ws3.append(
        [
            "序号", "水平档位", "采分点", "判定档位(0/1/2)", "得分",
            "多次判定一致率", "AI 理由", "踩点证据", "教师分", "AI总分",
        ]
    )
    for s in students:
        a = ai.get(s.seq)
        if not a:
            continue
        for pr in a.points:
            ws3.append(
                [
                    s.seq,
                    s.ability,
                    pr.text,
                    pr.level,
                    round(pr.score * pr.level / 2, 2),
                    pr.agreement,
                    pr.reason,
                    pr.evidence,
                    teacher.get(s.seq, ""),
                    a.total,
                ]
            )
    ws3.column_dimensions["C"].width = 40
    ws3.column_dimensions["G"].width = 40
    ws3.column_dimensions["H"].width = 40
    _style_header(ws3)

    wb.save(path)
    return path
